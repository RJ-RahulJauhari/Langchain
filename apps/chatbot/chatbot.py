# 4.Prompt/prompt_ui.py
from __future__ import annotations

import json
import shutil
import time
import uuid
import pathlib
from typing import Dict, List, Optional, Tuple

from dotenv import load_dotenv
load_dotenv()

import streamlit as st

# LangChain core & Ollama chat
from langchain_core.messages import HumanMessage, AIMessage, SystemMessage
from langchain_ollama import ChatOllama

# Embeddings + Vector Store (Chroma)
from langchain_community.embeddings import OllamaEmbeddings
from langchain_community.vectorstores import Chroma
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain.schema import Document

# Optional loaders
PDF_OK = True
DOCX_OK = True
try:
    from langchain_community.document_loaders import PyPDFLoader
except Exception:
    PDF_OK = False
try:
    import docx  # python-docx
except Exception:
    DOCX_OK = False


# =========================================================
# App Config
# =========================================================
st.set_page_config(page_title="Research Tool", layout="wide")

APP_DIR = pathlib.Path(".research_tool").resolve()
UPLOADS_DIR = APP_DIR / "uploads"
VSTORES_DIR = APP_DIR / "vstores"
STATE_FILE = APP_DIR / "chats.json"

DEFAULT_MODEL = "qwen3-coder:30b"
DEFAULT_TEMPERATURE = 0.4              # lower for factual work
DEFAULT_TOP_K = 12                     # more context chunks
DEFAULT_CHUNK_SIZE = 2000              # bigger chunks
DEFAULT_CHUNK_OVERLAP = 250

# Global, shared Chroma collection (persists across sessions & chats)
GLOBAL_COLLECTION = "global_docs"
GLOBAL_VSTORE_DIR = VSTORES_DIR / "global"

# Ensure dirs
APP_DIR.mkdir(exist_ok=True)
UPLOADS_DIR.mkdir(exist_ok=True)
VSTORES_DIR.mkdir(exist_ok=True)
GLOBAL_VSTORE_DIR.mkdir(parents=True, exist_ok=True)


# =========================================================
# Persistence (UI state)
# =========================================================
def save_state():
    data = {
        "chats": st.session_state.get("chats", {}),
        "active_chat_id": st.session_state.get("active_chat_id"),
        "settings": {
            "model_name": st.session_state.get("model_name", DEFAULT_MODEL),
            "temperature": st.session_state.get("temperature", DEFAULT_TEMPERATURE),
            "use_docs": st.session_state.get("use_docs", True),
            "use_memory": st.session_state.get("use_memory", True),
            "top_k": st.session_state.get("top_k", DEFAULT_TOP_K),
            "global_search": st.session_state.get("global_search", True),
            "retrieval_mode": st.session_state.get("retrieval_mode", "auto"),
            "confidence_threshold": st.session_state.get("confidence_threshold", 0.4),
            "chunk_size": st.session_state.get("chunk_size", DEFAULT_CHUNK_SIZE),
            "chunk_overlap": st.session_state.get("chunk_overlap", DEFAULT_CHUNK_OVERLAP),
            "use_mmr": st.session_state.get("use_mmr", False),
            "mmr_lambda": st.session_state.get("mmr_lambda", 0.5),
        },
    }
    STATE_FILE.write_text(json.dumps(data, ensure_ascii=False, indent=2))


def load_state():
    if STATE_FILE.exists():
        try:
            data = json.loads(STATE_FILE.read_text())
            st.session_state.chats = data.get("chats", {})
            st.session_state.active_chat_id = data.get("active_chat_id")
            s = data.get("settings", {})
            st.session_state.model_name = s.get("model_name", DEFAULT_MODEL)
            st.session_state.temperature = s.get("temperature", DEFAULT_TEMPERATURE)
            st.session_state.use_docs = s.get("use_docs", True)
            st.session_state.use_memory = s.get("use_memory", True)
            st.session_state.top_k = s.get("top_k", DEFAULT_TOP_K)
            st.session_state.global_search = s.get("global_search", True)
            st.session_state.retrieval_mode = s.get("retrieval_mode", "auto")
            st.session_state.confidence_threshold = s.get("confidence_threshold", 0.4)
            st.session_state.chunk_size = s.get("chunk_size", DEFAULT_CHUNK_SIZE)
            st.session_state.chunk_overlap = s.get("chunk_overlap", DEFAULT_CHUNK_OVERLAP)
            st.session_state.use_mmr = s.get("use_mmr", False)
            st.session_state.mmr_lambda = s.get("mmr_lambda", 0.5)
        except Exception:
            _reset_defaults()
    else:
        _reset_defaults()


def _reset_defaults():
    st.session_state.chats = {}
    st.session_state.active_chat_id = None
    st.session_state.model_name = DEFAULT_MODEL
    st.session_state.temperature = DEFAULT_TEMPERATURE
    st.session_state.use_docs = True
    st.session_state.use_memory = True
    st.session_state.top_k = DEFAULT_TOP_K
    st.session_state.global_search = True
    st.session_state.retrieval_mode = "auto"
    st.session_state.confidence_threshold = 0.4
    st.session_state.chunk_size = DEFAULT_CHUNK_SIZE
    st.session_state.chunk_overlap = DEFAULT_CHUNK_OVERLAP
    st.session_state.use_mmr = False
    st.session_state.mmr_lambda = 0.5


def ensure_state():
    if "chats" not in st.session_state:
        load_state()
    if not st.session_state.chats:
        new_chat()
    if (
        "active_chat_id" not in st.session_state
        or st.session_state.active_chat_id not in st.session_state.chats
    ):
        st.session_state.active_chat_id = next(iter(st.session_state.chats.keys()))
    save_state()


# =========================================================
# Chat management
# =========================================================
def new_chat(title: str = "New chat") -> str:
    cid = str(uuid.uuid4())[:8]
    st.session_state.chats[cid] = {
        "id": cid,
        "title": title,
        "created_at": time.time(),
        "messages": [],   # [{role, content, ts?, meta?}]
        "summary": "",    # rolling conversation summary
        "files": [],      # [{name, path}]
    }
    st.session_state.active_chat_id = cid
    (UPLOADS_DIR / cid).mkdir(parents=True, exist_ok=True)
    save_state()
    return cid


def chat_title_from(text: str) -> str:
    t = (text or "New chat").strip().replace("\n", " ")
    return (t[:40] + "…") if len(t) > 40 else t


def list_chats_sorted():
    return sorted(st.session_state.chats.items(), key=lambda x: x[1]["created_at"], reverse=True)


# =========================================================
# Model + Embeddings + Chroma (GLOBAL)
# =========================================================
@st.cache_resource
def get_model(name: str, temperature: float):
    return ChatOllama(model=name, temperature=temperature)

@st.cache_resource
def get_embedder():
    return OllamaEmbeddings(model="mxbai-embed-large")

@st.cache_resource
def get_global_vstore():
    embedder = get_embedder()
    return Chroma(
        collection_name=GLOBAL_COLLECTION,
        embedding_function=embedder,
        persist_directory=str(GLOBAL_VSTORE_DIR),
    )


def split_docs(texts_with_meta: List[Tuple[str, dict]], *, chunk_size: int, chunk_overlap: int):
    splitter = RecursiveCharacterTextSplitter(chunk_size=chunk_size, chunk_overlap=chunk_overlap)
    docs: List[Document] = []
    for text, meta in texts_with_meta:
        if not text:
            continue
        for chunk in splitter.split_text(text):
            m = {**meta, "type": meta.get("type", "doc_chunk")}
            docs.append(Document(page_content=chunk, metadata=m))
    return docs


# =========================================================
# File reading helpers
# =========================================================
def read_txt_like(path: pathlib.Path) -> str:
    try:
        return path.read_text(encoding="utf-8")
    except Exception:
        return path.read_text(encoding="latin-1", errors="ignore")


def read_docx(path: pathlib.Path) -> str:
    if not DOCX_OK:
        return ""
    d = docx.Document(str(path))
    return "\n".join(p.text for p in d.paragraphs)


def read_pdf(path: pathlib.Path) -> str:
    if not PDF_OK:
        return ""
    try:
        loader = PyPDFLoader(str(path))
        pages = loader.load()
        return "\n".join(p.page_content for p in pages)
    except Exception:
        return ""


def _build_texts_from_infos(cid: str, infos: List[dict]) -> List[Tuple[str, dict]]:
    texts_with_meta: List[Tuple[str, dict]] = []
    for info in infos:
        p = pathlib.Path(info["path"]) if isinstance(info, dict) else pathlib.Path(info)
        ext = p.suffix.lower()
        if ext in [".txt", ".md", ".markdown", ".csv", ".log"]:
            text = read_txt_like(p)
        elif ext == ".pdf":
            text = read_pdf(p)
        elif ext == ".docx":
            text = read_docx(p)
        else:
            try:
                text = p.read_text(encoding="utf-8")
            except Exception:
                text = ""
        if text:
            texts_with_meta.append((text, {"source": p.name, "chat_id": cid, "type": "doc_chunk"}))
    return texts_with_meta


# =========================================================
# Ingestion into GLOBAL Chroma (shared across chats)
# =========================================================
def ingest_files_global(cid: str, uploaded_files: List) -> List[dict]:
    chat = st.session_state.chats[cid]
    saved = []
    for f in uploaded_files or []:
        target_dir = UPLOADS_DIR / cid
        target_dir.mkdir(parents=True, exist_ok=True)
        dest = target_dir / f.name
        with open(dest, "wb") as w:
            w.write(f.read())
        info = {"name": f.name, "path": str(dest)}
        chat["files"].append(info)
        saved.append(info)

    if not saved:
        return saved

    texts_with_meta = _build_texts_from_infos(cid, saved)
    if not texts_with_meta:
        return saved

    vstore = get_global_vstore()
    docs = split_docs(
        texts_with_meta,
        chunk_size=int(st.session_state.get("chunk_size", DEFAULT_CHUNK_SIZE)),
        chunk_overlap=int(st.session_state.get("chunk_overlap", DEFAULT_CHUNK_OVERLAP)),
    )

    doc_ids = [f"doc-{d.metadata.get('chat_id','x')}-{uuid.uuid4()}" for d in docs]
    vstore.add_documents(documents=docs, ids=doc_ids)
    vstore.persist()
    save_state()
    return saved


def reindex_chat_docs(cid: str):
    """Delete this chat's existing doc chunks from Chroma and re-ingest using current chunk settings."""
    chat = st.session_state.chats[cid]
    vstore = get_global_vstore()
    col = vstore._collection  # underlying chromadb collection

    # --- delete existing vectors for this chat's document chunks ---
    deleted_ok = False
    try:
        # Chroma expects exactly one top-level operator
        col.delete(where={"$and": [{"type": "doc_chunk"}, {"chat_id": cid}]})
        deleted_ok = True
    except Exception:
        # Fallback: get IDs by filter then delete by ids
        try:
            got = col.get(where={"$and": [{"type": "doc_chunk"}, {"chat_id": cid}]}, include=["ids"])
            ids = got.get("ids", []) or []
            # sometimes 'ids' can be nested lists; flatten just in case
            if ids and isinstance(ids[0], list):
                ids = [i for sub in ids for i in sub]
            if ids:
                col.delete(ids=ids)
                deleted_ok = True
        except Exception as e:
            st.warning(f"Could not delete existing vectors (fallback failed): {e}")

    if not deleted_ok:
        st.info("Proceeding without prior-vector deletion (duplicates may remain).")

    # --- rebuild with current chunking settings ---
    texts_with_meta = _build_texts_from_infos(cid, chat.get("files", []))
    if not texts_with_meta:
        st.info("No files found to re-index for this chat.")
        return 0

    docs = split_docs(
        texts_with_meta,
        chunk_size=int(st.session_state.get("chunk_size", 2000)),
        chunk_overlap=int(st.session_state.get("chunk_overlap", 250)),
    )
    vstore.add_documents(docs, ids=[f"doc-{cid}-{uuid.uuid4()}" for _ in docs])
    vstore.persist()
    return len(docs)



# =========================================================
# Index chat messages into GLOBAL Chroma
# =========================================================
def index_chat_message(cid: str, role: str, content: str):
    vstore = get_global_vstore()
    chat = st.session_state.chats[cid]
    meta = {
        "type": "chat_message",
        "role": role,
        "chat_id": cid,
        "chat_title": chat["title"],
        "ts": time.time(),
        "source": "chat_history",
    }
    doc = Document(page_content=content, metadata=meta)
    vstore.add_documents([doc], ids=[f"msg-{cid}-{uuid.uuid4()}"])
    vstore.persist()


# =========================================================
# Retrieval (GLOBAL, docs only)
# =========================================================
def retrieve_docs(query: str, cid: str, top_k: int, global_search: bool) -> List[dict]:
    vstore = get_global_vstore()
    filters = {"type": "doc_chunk"}
    if not global_search:
        filters["chat_id"] = cid

    use_mmr = bool(st.session_state.get("use_mmr", False))
    mmr_lambda = float(st.session_state.get("mmr_lambda", 0.5))

    results = []
    if use_mmr:
        try:
            docs = vstore.max_marginal_relevance_search(
                query,
                k=top_k,
                fetch_k=max(top_k * 5, top_k + 10),
                filter=filters,
                lambda_mult=mmr_lambda,
            )
            results = [(d, None) for d in docs]
        except Exception:
            pass

    if not results:
        try:
            results = vstore.similarity_search_with_score(query, k=top_k, filter=filters)
        except TypeError:
            docs = vstore.similarity_search(query, k=top_k, filter=filters)
            results = [(d, None) for d in docs]

    if not results and global_search:
        try:
            results = vstore.similarity_search_with_score(query, k=top_k)
        except TypeError:
            docs = vstore.similarity_search(query, k=top_k)
            results = [(d, None) for d in docs]

    # Deduplicate identical contents to reduce repeats
    seen = set()
    contexts = []
    for d, score in results:
        key = hash(d.page_content.strip()[:500])
        if key in seen:
            continue
        seen.add(key)
        contexts.append({
            "content": d.page_content,
            "source": d.metadata.get("source", "N/A"),
            "chat_id": d.metadata.get("chat_id"),
            "type": d.metadata.get("type"),
            "score": float(score) if score is not None else None,
        })
    return contexts


# =========================================================
# History helpers
# =========================================================
def is_history_query(q: str) -> bool:
    ql = q.lower().strip()
    triggers = [
        "what did i ask previously",
        "what did i ask before",
        "what did we discuss before",
        "what did we talk about before",
        "previous chats",
        "show my past questions",
        "what did i ask last time",
        "remind me what i asked",
    ]
    return any(t in ql for t in triggers)


def search_history_similarity(query: str, cid: str, top_k: int, global_search: bool, role: Optional[str] = "user"):
    vstore = get_global_vstore()
    filters = {"type": "chat_message"}
    if role:
        filters["role"] = role
    if not global_search:
        filters["chat_id"] = cid
    return vstore.similarity_search(query if query.strip() else "recent", k=top_k, filter=filters)


def collect_recent_user_msgs_from_state(limit: int = 10, cid: Optional[str] = None):
    items = []
    for chat_id, chat in st.session_state.chats.items():
        if cid and chat_id != cid:
            continue
        for m in chat.get("messages", []):
            if m.get("role") == "user":
                items.append({
                    "chat_id": chat_id,
                    "chat_title": chat.get("title", "(untitled)"),
                    "content": m.get("content", ""),
                    "ts": m.get("ts", 0),
                })
    items.sort(key=lambda x: x["ts"], reverse=True)
    return items[:limit]


# =========================================================
# Auto RAG: LLM decides whether to retrieve
# =========================================================
def _extract_json(text: str) -> dict:
    try:
        start = text.find("{")
        end = text.rfind("}")
        if start == -1 or end == -1 or end <= start:
            return {}
        return json.loads(text[start:end+1])
    except Exception:
        return {}


def decide_retrieval_llm(model: ChatOllama, question: str, summary: str, recent: List[Dict[str, str]], default_top_k: int) -> dict:
    recent_text = "\n".join(f"- {m['role']}: {m['content'][:200]}" for m in recent[-6:])
    prompt = f"""
You are a retrieval controller. Decide if the user question needs external documents (RAG) or not.
Respond with ONLY a compact JSON object of the form:
{{"need_retrieval": true|false, "confidence": 0.0-1.0, "reason": "...", "suggested_query": "...", "top_k": 1-10}}

Question: {question}
Conversation summary (may be empty): {summary or '(none)'}
Recent messages:
{recent_text if recent_text else '(none)'}

Heuristics: If the question asks for facts from uploaded files, citations, statistics, dates, figures, or specific sections, prefer retrieval.
If it's a general explanation or small coding tip that does not depend on the user's documents, no retrieval.
""".strip()
    try:
        resp = model.invoke([HumanMessage(content=prompt)])
        data = _extract_json(getattr(resp, "content", ""))
    except Exception:
        data = {}

    if not isinstance(data, dict) or "need_retrieval" not in data:
        low = question.lower()
        needles = ["according to", "in the pdf", "in the doc", "reference", "cite", "citation", "paper", "report", "dataset", "from the file"]
        need = any(k in low for k in needles) or len(question) > 140
        data = {
            "need_retrieval": need,
            "confidence": 0.55 if need else 0.35,
            "reason": "heuristic fallback",
            "suggested_query": question,
            "top_k": default_top_k,
        }
    else:
        data.setdefault("top_k", default_top_k)
        data.setdefault("suggested_query", question)
        data.setdefault("reason", "")
        data.setdefault("confidence", 0.5)
        data["need_retrieval"] = bool(data.get("need_retrieval"))
    return data


# =========================================================
# Prompt assembly + memory compression
# =========================================================
def build_system_prompt() -> str:
    # Strict: use Provided Context and never disclaim lack of access if context exists
    return (
        "ROLE: You are a focused research assistant with retrieval.\n"
        "WHEN CONTEXT IS PROVIDED:\n"
        "  - You DO have access to the 'Provided Context' section and must rely on it.\n"
        "  - Never say you lack access to documents or prior files when context is provided.\n"
        "  - Prefer citing those sources inline as **[source: filename]**.\n"
        "IF INFORMATION IS MISSING:\n"
        "  - Say: 'Not in provided context.' and list what else is needed.\n"
        "FORBIDDEN PHRASES (unless context is empty): 'I don't have access', 'I cannot access previous chats', 'I don't have the document'.\n"
        "OUTPUT: Clean Markdown with a concise title, headings, bullets, and brief code blocks only if helpful."
    )


def summarize_history_if_needed(model: ChatOllama, chat: dict, max_chars: int = 4000) -> None:
    history_text = "\n\n".join(f"{m['role']}: {m['content']}" for m in chat["messages"])
    if len(history_text) < max_chars:
        return
    try:
        with st.spinner("Compressing history…"):
            prompt = (
                "Summarize the following conversation so far into a compact brief that preserves:\n"
                "- main questions, assumptions, decisions, and key facts\n"
                "- do NOT add new information\n\n"
                f"Conversation:\n{history_text}"
            )
            resp = model.invoke([HumanMessage(content=prompt)])
            summary = getattr(resp, "content", "")
            recent = chat["messages"][-8:] if len(chat["messages"]) > 8 else chat["messages"]
            chat["messages"] = [{"role": "system", "content": f"Summary so far: {summary}", "ts": time.time()}] + recent
            chat["summary"] = summary
    except Exception as e:
        st.warning(f"History summarization failed (continuing without compression): {e}")


def to_lc_messages(history: List[Dict[str, str]], user_input: str, context_blocks: List[dict]) -> List:
    """Build the message list for the model.
    IMPORTANT: Provided Context is injected as a HumanMessage so it is never ignored by backends that only use the first SystemMessage.
    """
    msgs: List = []
    # 1) System rules
    msgs.append(SystemMessage(content=build_system_prompt()))

    # 2) Conversation history (trimmed/summary already handled upstream)
    for m in history:
        if m["role"] == "user":
            msgs.append(HumanMessage(content=m["content"]))
        elif m["role"] == "assistant":
            msgs.append(AIMessage(content=m["content"]))
        elif m["role"] == "system":
            msgs.append(SystemMessage(content=m["content"]))

    # 3) Provided Context as HUMAN message (so it's definitely passed)
    if context_blocks:
        parts = []
        for i, b in enumerate(context_blocks, 1):
            src = b.get("source", "N/A")
            snippet = (b.get("content", "") or "").strip()
            if len(snippet) > 1600:
                snippet = snippet[:1600] + "…"
            parts.append(f"### Source {i} — {src}\n```\n{snippet}\n```")
        ctx_text = "## PROVIDED CONTEXT\n" + "\n\n".join(parts) + "\n\nUse ONLY this context for facts and cite like **[source: filename]**."
        msgs.append(HumanMessage(content=ctx_text))

    # 4) Final guarded user request
    user_guard = (
        "INSTRUCTIONS:\n"
        " - If 'PROVIDED CONTEXT' is present above, use it and do NOT say you lack access.\n"
        " - Cite sources like **[source: filename]**.\n"
        " - If an answer is not supported by the context, write: 'Not in provided context.'\n\n"
        f"QUESTION:\n{user_input}"
    )
    msgs.append(HumanMessage(content=user_guard))
    return msgs


# =========================================================
# Sidebar UI
# =========================================================
def sidebar_ui():
    with st.sidebar:
        st.title("💬 Chats")

        if st.button("➕ New chat", use_container_width=True):
            new_chat()

        st.divider()
        for cid, chat in list_chats_sorted():
            label = chat["title"]
            is_active = (cid == st.session_state.active_chat_id)
            btn_label = f"▶ {label}" if is_active else f"• {label}"
            if st.button(btn_label, key=f"sel_{cid}", use_container_width=True):
                st.session_state.active_chat_id = cid
                save_state()

        st.divider()
        active = st.session_state.chats[st.session_state.active_chat_id]
        new_title = st.text_input("Rename this chat", value=active["title"], key="rename_input")
        if new_title and new_title != active["title"]:
            active["title"] = new_title
            save_state()

        c1, c2 = st.columns(2)
        with c1:
            if st.button("🧹 Clear chat"):
                active["messages"].clear()
                active["summary"] = ""
                save_state()
        with c2:
            if st.button("🗑 Delete chat"):
                try:
                    shutil.rmtree(UPLOADS_DIR / active["id"], ignore_errors=True)
                except Exception:
                    pass
                del st.session_state.chats[active["id"]]
                if not st.session_state.chats:
                    new_chat()
                st.session_state.active_chat_id = next(iter(st.session_state.chats.keys()))
                save_state()
                st.experimental_rerun()

        st.divider()
        st.subheader("Model")
        st.session_state.model_name = st.text_input("Model name", value=st.session_state.get("model_name", DEFAULT_MODEL))
        st.session_state.temperature = st.slider("Temperature", 0.0, 2.0, float(st.session_state.get("temperature", DEFAULT_TEMPERATURE)), 0.1)

        st.divider()
        st.subheader("Documents & Memory")
        st.session_state.use_docs = st.toggle("Use uploaded documents (Chroma)", value=bool(st.session_state.get("use_docs", True)))
        st.session_state.global_search = st.toggle("Search across all chats (global RAG)", value=bool(st.session_state.get("global_search", True)))
        st.session_state.top_k = st.slider("Top-K (context chunks)", 1, 50, int(st.session_state.get("top_k", DEFAULT_TOP_K)))
        st.session_state.use_memory = st.toggle("Use conversation memory", value=bool(st.session_state.get("use_memory", True)))

        st.caption("Chunking (affects NEW ingests; use re-index for old files)")
        st.session_state.chunk_size = st.number_input("Chunk size", min_value=300, max_value=4000, step=100, value=int(st.session_state.get("chunk_size", DEFAULT_CHUNK_SIZE)))
        st.session_state.chunk_overlap = st.number_input("Chunk overlap", min_value=0, max_value=1000, step=50, value=int(st.session_state.get("chunk_overlap", DEFAULT_CHUNK_OVERLAP)))
        if st.session_state.chunk_overlap >= st.session_state.chunk_size:
            st.warning("Chunk overlap must be smaller than chunk size; adjusting…")
            st.session_state.chunk_overlap = max(0, st.session_state.chunk_size // 4)

        st.session_state.use_mmr = st.toggle("Use MMR (diverse chunks)", value=bool(st.session_state.get("use_mmr", False)))
        if st.session_state.use_mmr:
            st.session_state.mmr_lambda = st.slider("MMR diversity (λ)", 0.0, 1.0, float(st.session_state.get("mmr_lambda", 0.5)), 0.05)

        # Retrieval mode controls
        mode_label = {"auto": "Auto (LLM)", "always": "Always", "never": "Never"}[st.session_state.get("retrieval_mode", "auto")]
        mode = st.selectbox("Retrieval mode", ["Auto (LLM)", "Always", "Never"], index=["Auto (LLM)", "Always", "Never"].index(mode_label))
        st.session_state.retrieval_mode = {"Auto (LLM)": "auto", "Always": "always", "Never": "never"}[mode]
        if st.session_state.retrieval_mode == "auto":
            st.session_state.confidence_threshold = st.slider("Auto mode: confidence threshold", 0.0, 1.0, float(st.session_state.get("confidence_threshold", 0.4)), 0.05)

        # Upload
        st.markdown("**Upload files** (txt, md, pdf, docx, csv, log)")
        files = st.file_uploader(
            "Add to the global Chroma store (tagged with this chat)",
            type=["txt", "md", "markdown", "pdf", "docx", "csv", "log"],
            accept_multiple_files=True,
            label_visibility="collapsed",
        )
        if files:
            with st.spinner("Indexing into Chroma (global)…"):
                added = ingest_files_global(st.session_state.active_chat_id, files)
                st.success(f"Added {len(added)} file(s) to global store")
                save_state()

        # Re-index button
        if st.button("🔁 Re-index this chat's files with current chunking"):
            with st.spinner("Re-indexing…"):
                n = reindex_chat_docs(st.session_state.active_chat_id)
                st.success(f"Re-indexed {n} chunks for this chat.")
                save_state()

        # List current chat files
        active_files = st.session_state.chats[st.session_state.active_chat_id].get("files", [])
        if active_files:
            st.markdown("**Files uploaded in this chat** (now searchable globally)")
            for info in active_files:
                st.caption(f"• {info['name']}")

        st.divider()
        st.subheader("🔎 Search history")
        q = st.text_input("Find past questions/answers", key="history_query")
        if st.button("Search history"):
            if q.strip():
                docs = search_history_similarity(q, st.session_state.active_chat_id, top_k=10, global_search=st.session_state.global_search, role=None)
            else:
                recents = collect_recent_user_msgs_from_state(limit=10, cid=None if st.session_state.global_search else st.session_state.active_chat_id)
                st.session_state["_history_sidebar_results"] = [
                    {
                        "chat_title": r["chat_title"],
                        "snippet": (r["content"][:160] + "…") if len(r["content"]) > 160 else r["content"],
                        "ts": r["ts"],
                        "chat_id": r["chat_id"],
                    }
                    for r in recents
                ]
                docs = None

            if docs is not None:
                results = []
                for d in docs:
                    md = d.metadata or {}
                    results.append({
                        "chat_title": md.get("chat_title", "(untitled)"),
                        "snippet": (d.page_content[:160] + "…") if len(d.page_content) > 160 else d.page_content,
                        "ts": md.get("ts", 0),
                        "chat_id": md.get("chat_id"),
                        "role": md.get("role"),
                    })
                st.session_state["_history_sidebar_results"] = results

        results = st.session_state.get("_history_sidebar_results")
        if results:
            for r in results:
                ts_str = time.strftime('%Y-%m-%d %H:%M', time.localtime(r.get("ts", 0))) if r.get("ts") else ""
                st.caption(f"[{ts_str}] {r['chat_title']} — {r['snippet']}")


# =========================================================
# Main UI
# =========================================================
def render_messages(msgs: List[Dict[str, str]]):
    for m in msgs:
        with st.chat_message("user" if m["role"] == "user" else ("assistant" if m["role"] == "assistant" else "system")):
            if m.get("meta") and m["role"] == "assistant":
                meta = m["meta"]
                if meta.get("used_rag"):
                    scope = meta.get("scope", "global")
                    k = meta.get("k")
                    st.caption(f"📚 Used Chroma (scope: {scope}, k={k})")
                else:
                    st.caption("💡 No retrieval — answered from memory/knowledge")
            st.markdown(m["content"])


def handle_history_query(user_msg: str, active: dict):
    q = user_msg.strip()

    lowers = q.lower()
    topic = ""
    for prefix in [
        "what did i ask previously about",
        "what did i ask before about",
        "what did we discuss before about",
        "remind me what i asked about",
    ]:
        if lowers.startswith(prefix):
            topic = q[len(prefix):].strip(" :?.,")
            break

    global_search = bool(st.session_state.get("global_search", True))

    if not topic:
        recents = collect_recent_user_msgs_from_state(limit=12, cid=None if global_search else active["id"])
        if not recents:
            st.markdown("_No previous questions found._")
            return
        lines = ["### Recent things you asked\n"]
        for r in recents:
            ts_str = time.strftime('%Y-%m-%d %H:%M', time.localtime(r.get("ts", 0))) if r.get("ts") else ""
            title = r.get("chat_title", "(untitled)")
            snippet = r.get("content", "").strip()
            if len(snippet) > 180:
                snippet = snippet[:180] + "…"
            lines.append(f"- **[{ts_str}]** *{title}* — {snippet} **[chat: {r['chat_id']}]**")
        st.markdown("\n".join(lines))
        st.markdown("_Tip: ask ‘what did I ask previously about <topic>?’, and I’ll pull exact matches with references._")
        return

    docs = search_history_similarity(topic, active["id"], top_k=12, global_search=global_search, role="user")
    if not docs:
        st.markdown("_Couldn't find past questions on that topic._")
        return

    st.markdown(f"### Past questions about **{topic}**")
    for d in docs:
        md = d.metadata or {}
        ts = md.get("ts", 0)
        ts_str = time.strftime('%Y-%m-%d %H:%M', time.localtime(ts)) if ts else ""
        title = md.get("chat_title", "(untitled)")
        snippet = d.page_content.strip()
        if len(snippet) > 220:
            snippet = snippet[:220] + "…"
        st.markdown(f"- **[{ts_str}]** *{title}* — {snippet} **[chat: {md.get('chat_id')}]**")


def main():
    ensure_state()
    sidebar_ui()

    st.header("🔎 Research Tool")

    active = st.session_state.chats[st.session_state.active_chat_id]
    render_messages([m for m in active["messages"] if m["role"] != "system"])

    user_msg = st.chat_input("Ask a research question, or say ‘what did I ask previously’ …")
    if user_msg:
        # Add user message (+ timestamp)
        active["messages"].append({"role": "user", "content": user_msg, "ts": time.time()})
        if active["title"] == "New chat":
            active["title"] = chat_title_from(user_msg)
        save_state()

        # Index this user message into global Chroma history
        try:
            index_chat_message(active["id"], role="user", content=user_msg)
        except Exception as e:
            st.warning(f"History indexing failed (user msg): {e}")

        # Show immediately
        with st.chat_message("user"):
            st.markdown(user_msg)

        with st.chat_message("assistant"):
            # 1) Intercept special history queries
            if is_history_query(user_msg):
                handle_history_query(user_msg, active)
                return

            # 2) Normal RAG flow
            with st.spinner("Thinking…"):
                try:
                    model = get_model(st.session_state.model_name, st.session_state.temperature)

                    if st.session_state.use_memory:
                        summarize_history_if_needed(model, active, max_chars=4000)

                    contexts: List[dict] = []
                    used_rag = False
                    retrieved_k = None
                    scope_label = "global" if st.session_state.get("global_search", True) else "chat"
                    retrieval_attempted = False

                    # Decide retrieval per mode
                    retrieval_mode = st.session_state.get("retrieval_mode", "auto")
                    if st.session_state.get("use_docs", True) and retrieval_mode != "never":
                        if retrieval_mode == "always":
                            k = int(st.session_state.top_k)   # from slider
                            contexts = retrieve_docs(
                                user_msg,
                                active["id"],
                                k,
                                st.session_state.global_search,
                            )
                            retrieval_attempted = True
                            used_rag = len(contexts) > 0
                            retrieved_k = len(contexts)
                        else:  # auto via LLM
                            decision = decide_retrieval_llm(
                                model,
                                question=user_msg,
                                summary=active.get("summary", ""),
                                recent=active["messages"][-6:],
                                default_top_k=st.session_state.top_k,
                            )
                            if decision.get("need_retrieval") and float(decision.get("confidence", 0)) >= float(st.session_state.get("confidence_threshold", 0.4)):
                                q = decision.get("suggested_query") or user_msg
                                k = int(st.session_state.top_k)  # always use slider K
                                contexts = retrieve_docs(q, active["id"], k, st.session_state.global_search)
                                retrieval_attempted = True
                                used_rag = len(contexts) > 0
                                retrieved_k = len(contexts)

                    # Indicator of retrieval usage
                    if used_rag:
                        st.caption(f"📚 Used Chroma (scope: {scope_label}, k={retrieved_k})")
                    else:
                        st.caption("💡 No retrieval — answered from memory/knowledge")
                        if retrieval_attempted and not contexts:
                            st.caption("🔎 No relevant chunks found in Chroma for this query.")

                    # Show retrieved chunks
                    if used_rag and contexts:
                        with st.expander(f"Show retrieved context (k={len(contexts)})", expanded=False):
                            for i, c in enumerate(contexts, 1):
                                score = c.get("score")
                                score_txt = f" • distance: {score:.4f}" if isinstance(score, float) else ""
                                src = c.get("source", "N/A")
                                cid_ref = c.get("chat_id", "-")
                                st.markdown(f"**Chunk {i}** — *{src}* (chat: `{cid_ref}`){score_txt}")
                                snippet = c["content"].strip()
                                if len(snippet) > 1600:
                                    snippet = snippet[:1600] + "…"
                                st.code(snippet)
                                st.divider()

                    # >>> THIS IS THE FIX: the Provided Context is now inside the message list as a HumanMessage
                    lc_msgs = to_lc_messages(active["messages"], user_msg, contexts)
                    resp = model.invoke(lc_msgs)
                    answer = getattr(resp, "content", str(resp))

                    st.markdown(answer)
                    # Save assistant reply (+ ts + meta)
                    active["messages"].append({
                        "role": "assistant",
                        "content": answer,
                        "ts": time.time(),
                        "meta": {"used_rag": used_rag, "k": retrieved_k, "scope": scope_label},
                    })
                    save_state()

                    # Index assistant reply into global history
                    try:
                        index_chat_message(active["id"], role="assistant", content=answer)
                    except Exception as e:
                        st.warning(f"History indexing failed (assistant msg): {e}")

                except Exception as e:
                    st.error(f"Model error: {e}")


if __name__ == "__main__":
    main()
